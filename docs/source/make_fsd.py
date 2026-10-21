"""Restaurant OS — Functional Specification Document (FSD) deck with DFD Level 0 and Level 1, as native PowerPoint shapes.

    python3 docs/source/make_fsd.py docs/Restaurant_OS_FSD.pptx   (needs python-pptx)
"""

import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

OUT = sys.argv[1] if len(sys.argv) > 1 else "Restaurant_OS_FSD.pptx"

# palette
TEAL, TEAL_D, TEAL_L = RGBColor(0x0D, 0x94, 0x88), RGBColor(0x0B, 0x5E, 0x57), RGBColor(0xE6, 0xF6, 0xF4)
INK, MUTED, LINE = RGBColor(0x1E, 0x29, 0x3B), RGBColor(0x64, 0x74, 0x8B), RGBColor(0xCB, 0xD5, 0xE1)
ORANGE, ORANGE_L = RGBColor(0xEA, 0x68, 0x34), RGBColor(0xFD, 0xF0, 0xE9)
BLUE, BLUE_L = RGBColor(0x2A, 0x78, 0xD6), RGBColor(0xEA, 0xF2, 0xFC)
WHITE, BG = RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xF8, 0xFA, 0xFC)
FONT = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = prs.slide_width, prs.slide_height
state = {"n": 0}
i_ = Inches


# ------------------------------------------------------------------ primitives
def _font(run, size, bold=False, color=INK, italic=False):
	run.font.name, run.font.size, run.font.bold, run.font.italic = FONT, Pt(size), bold, italic
	run.font.color.rgb = color


def text(slide, x, y, w, h, s, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
	tb = slide.shapes.add_textbox(x, y, w, h)
	tf = tb.text_frame
	tf.word_wrap = True
	tf.margin_left = tf.margin_right = Inches(0.05)
	tf.margin_top = tf.margin_bottom = Inches(0.02)
	tf.vertical_anchor = anchor
	for i, line in enumerate(s.split("\n")):
		p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
		p.alignment = align
		_font(p.add_run(), size, bold, color, italic)
		p.runs[0].text = line
	return tb


def bullets(slide, x, y, w, h, items, size=14, color=INK, gap=4):
	"""items: str or (str, level). A '**Lead:** rest' prefix renders the lead in bold."""
	tb = slide.shapes.add_textbox(x, y, w, h)
	tf = tb.text_frame
	tf.word_wrap = True
	for i, it in enumerate(items):
		s, lvl = (it, 0) if isinstance(it, str) else it
		p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
		p.space_after = Pt(gap)
		p.level = lvl
		mark = "•  " if lvl == 0 else "–  "
		r0 = p.add_run(); r0.text = ("      " * lvl) + mark; _font(r0, size, color=TEAL if lvl == 0 else MUTED, bold=True)
		if s.startswith("**") and "**" in s[2:]:
			lead, rest = s[2:].split("**", 1)
			r = p.add_run(); r.text = lead; _font(r, size - (1 if lvl else 0), True, color)
			r = p.add_run(); r.text = rest; _font(r, size - (1 if lvl else 0), False, color)
		else:
			r = p.add_run(); r.text = s; _font(r, size - (1 if lvl else 0), False, color)
	return tb


def box(slide, x, y, w, h, s="", fill=WHITE, line=LINE, shape=MSO_SHAPE.RECTANGLE, size=12, bold=False, color=INK,
		align=PP_ALIGN.CENTER, lw=1.0, shadow=False):
	sh = slide.shapes.add_shape(shape, x, y, w, h)
	sh.fill.solid(); sh.fill.fore_color.rgb = fill
	if line is None:
		sh.line.fill.background()
	else:
		sh.line.color.rgb = line; sh.line.width = Pt(lw)
	if not shadow:
		style = sh._element.find(qn("p:style"))
		if style is not None:
			sh._element.remove(style)                         # no theme shadow / glow
	tf = sh.text_frame
	tf.word_wrap = True
	tf.margin_left = tf.margin_right = Inches(0.06)
	tf.margin_top = tf.margin_bottom = Inches(0.03)
	tf.vertical_anchor = MSO_ANCHOR.MIDDLE
	if s:
		for i, line_ in enumerate(s.split("\n")):
			p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
			p.alignment = align
			r = p.add_run(); r.text = line_
			_font(r, size if i == 0 else size - 2, bold if i == 0 else False, color if i == 0 else MUTED)
	return sh


def arrow(slide, x1, y1, x2, y2, label=None, color=MUTED, lw=1.25, size=9, lx=None, ly=None, both=False, dash=False):
	c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
	st_ = c._element.find(qn("p:style"))
	if st_ is not None:
		c._element.remove(st_)
	c.line.color.rgb = color; c.line.width = Pt(lw)
	ln = c.line._get_or_add_ln()
	if dash:
		ln.append(ln.makeelement(qn("a:prstDash"), {"val": "dash"}))
	ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
	if both:
		ln.append(ln.makeelement(qn("a:headEnd"), {"type": "triangle", "w": "med", "len": "med"}))
	if label:
		mx = lx if lx is not None else (x1 + x2) // 2 - Inches(0.9)
		my = ly if ly is not None else (y1 + y2) // 2 - Inches(0.14)
		t = text(slide, mx, my, Inches(1.8), Inches(0.3), label, size=size, color=INK, align=PP_ALIGN.CENTER)
		t.fill.solid(); t.fill.fore_color.rgb = WHITE
	return c


def table(slide, x, y, w, rows, col_w, size=11, head_fill=TEAL, row_h=0.36):
	n_r, n_c = len(rows), len(rows[0])
	gt = slide.shapes.add_table(n_r, n_c, x, y, w, Inches(row_h * n_r)).table
	for j, cw in enumerate(col_w):
		gt.columns[j].width = Inches(cw)
	for i, row in enumerate(rows):
		gt.rows[i].height = Inches(row_h)
		for j, val in enumerate(row):
			cell = gt.cell(i, j)
			cell.margin_left = cell.margin_right = Inches(0.07)
			cell.margin_top = cell.margin_bottom = Inches(0.03)
			cell.vertical_anchor = MSO_ANCHOR.MIDDLE
			cell.fill.solid()
			cell.fill.fore_color.rgb = head_fill if i == 0 else (WHITE if i % 2 else BG)
			tf = cell.text_frame
			tf.word_wrap = True
			tf.paragraphs[0].text = ""
			r = tf.paragraphs[0].add_run(); r.text = str(val)
			_font(r, size, bold=(i == 0 or j == 0), color=WHITE if i == 0 else INK)
	return gt


def slide(title, kicker=None, number=True):
	s = prs.slides.add_slide(BLANK)
	bg = s.background.fill; bg.solid(); bg.fore_color.rgb = WHITE
	box(s, 0, 0, W, Inches(0.09), fill=TEAL, line=None)
	if kicker:
		text(s, Inches(0.55), Inches(0.28), Inches(9), Inches(0.3), kicker.upper(), size=11, bold=True, color=TEAL)
	text(s, Inches(0.55), Inches(0.52 if kicker else 0.35), Inches(12.2), Inches(0.7), title, size=26, bold=True, color=INK)
	state["n"] += 1
	if number:
		text(s, Inches(0.55), H - Inches(0.42), Inches(8), Inches(0.3), "Restaurant OS · Functional Specification · v1.0 · 21 Oct 2026",
			 size=9, color=MUTED)
		text(s, W - Inches(1.3), H - Inches(0.42), Inches(0.8), Inches(0.3), str(state["n"]), size=9, color=MUTED, align=PP_ALIGN.RIGHT)
	return s


def notes(s, txt):
	s.notes_slide.notes_text_frame.text = txt


# ------------------------------------------------------------------ DFD shapes (Yourdon/DeMarco-style)
def entity(s, x, y, w, h, label, sub=None):
	return box(s, x, y, w, h, label + (f"\n{sub}" if sub else ""), fill=BLUE_L, line=BLUE, size=12, bold=True, lw=1.5)


def process(s, x, y, w, h, num, label, big=False):
	sh = box(s, x, y, w, h, "", fill=TEAL_L, line=TEAL_D, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=1.75)
	tf = sh.text_frame
	p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
	r_ = p.add_run(); r_.text = num; _font(r_, 10, True, TEAL)
	p = tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER
	r_ = p.add_run(); r_.text = label; _font(r_, 13, True, TEAL_D)
	return sh


def store(s, x, y, w, h, code, label):
	"""Open-ended data store: a box with top + bottom rules and a code cell on the left."""
	sh = box(s, x, y, w, h, fill=ORANGE_L, line=None)
	box(s, x, y, Inches(0.5), h, code, fill=ORANGE, line=None, size=11, bold=True, color=WHITE)
	text(s, x + Inches(0.55), y, w - Inches(0.55), h, label, size=11, bold=True, color=INK, anchor=MSO_ANCHOR.MIDDLE)
	for yy in (y, y + h):
		ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x, yy, x + w, yy)
		ln.line.color.rgb = ORANGE; ln.line.width = Pt(1.5)
	return sh


def cx(sh):
	return sh.left + sh.width // 2


def cy(sh):
	return sh.top + sh.height // 2


# ================================================================== slides
# 1. title
s = prs.slides.add_slide(BLANK)
s.background.fill.solid(); s.background.fill.fore_color.rgb = TEAL_D
box(s, 0, Inches(5.9), W, Inches(1.6), fill=TEAL, line=None)
text(s, Inches(0.8), Inches(1.4), Inches(11), Inches(0.5), "FUNCTIONAL SPECIFICATION DOCUMENT", size=14, bold=True, color=RGBColor(0x99, 0xF6, 0xE4))
text(s, Inches(0.8), Inches(1.9), Inches(11.5), Inches(1.2), "Restaurant OS", size=54, bold=True, color=WHITE)
text(s, Inches(0.8), Inches(3.0), Inches(11.5), Inches(1.0), "Automate the restaurant from one screen — orders, kitchen, riders, stock,\nreviews and reports on ERPNext v16",
	 size=20, color=RGBColor(0xD1, 0xFA, 0xF5))
text(s, Inches(0.8), Inches(4.5), Inches(11.5), Inches(0.9), "Includes Data Flow Diagrams — Level 0 (context) and Level 1",
	 size=16, bold=True, color=WHITE)
text(s, Inches(0.8), Inches(6.2), Inches(8), Inches(0.9), "Version 1.0  ·  21 October 2026\nDemo company: Saffron Grill (SGR) · Main branch: Gulberg branch at D-Ground, Faisalabad",
	 size=14, color=WHITE)
text(s, Inches(8.8), Inches(6.2), Inches(3.8), Inches(0.9), "App: restaurant_os\nSite: restaurant.localhost", size=14, color=WHITE, align=PP_ALIGN.RIGHT)
state["n"] += 1

# 2. document control + contents
s = slide("Document control & contents", "About this document")
table(s, Inches(0.55), Inches(1.35), Inches(5.6), [
	["Item", "Detail"],
	["Document", "Functional Specification (FSD)"],
	["Version", "1.0 — all 11 build phases complete"],
	["Date", "21 October 2026"],
	["System", "Restaurant OS (Frappe app restaurant_os)"],
	["Platform", "ERPNext v16 · Frappe v16 · React 18"],
	["Status", "Built, tested (37 automated tests)"],
	["Audience", "Owner, managers, implementers, developers"],
], [1.5, 4.1], size=12, row_h=0.42)
bullets(s, Inches(6.7), Inches(1.35), Inches(6.1), Inches(5.5), [
	"**1  ** Purpose, scope and users",
	"**2  ** System overview and the nine-step flow",
	"**3  ** Data Flow Diagram — Level 0 (context)",
	"**4  ** Data Flow Diagram — Level 1 (processes and data stores)",
	"**5  ** Functional requirements by module",
	"**6  ** Order state machine and business rules",
	"**7  ** Integrations, security and audit",
	"**8  ** Non-functional requirements and measured results",
	"**9  ** Acceptance criteria and open items",
], size=15, gap=8)

# 3. purpose & scope
s = slide("Purpose and scope", "1 · Introduction")
text(s, Inches(0.55), Inches(1.3), Inches(12.2), Inches(0.8),
	 "This document specifies what Restaurant OS does: the functions each user gets, the data that moves between them, and the rules the "
	 "system enforces. It describes the system as built and tested.", size=15, color=MUTED)
box(s, Inches(0.55), Inches(2.25), Inches(6.0), Inches(0.45), "In scope", fill=TEAL, line=None, size=14, bold=True, color=WHITE)
bullets(s, Inches(0.55), Inches(2.8), Inches(6.0), Inches(4.2), [
	"Orders from POS counter, dine-in tables, website / app, self-order kiosk, phone and Foodpanda — one queue",
	"Kitchen Display (KDS) by station; table map for dine-in",
	"Own riders and Foodpanda riders; live tracking; cash on delivery",
	"POS shifts, payments (cash, card, bank, online), FBR invoice reporting",
	"Recipe-based stock deduction, wastage, variance, auto-reorder",
	"Offers, coupons, loyalty points, reviews, campaigns",
	"Owner dashboard and reports",
], size=13)
box(s, Inches(6.85), Inches(2.25), Inches(5.95), Inches(0.45), "Out of scope (this version)", fill=MUTED, line=None, size=14, bold=True, color=WHITE)
bullets(s, Inches(6.85), Inches(2.8), Inches(5.95), Inches(4.2), [
	"Payroll and HR beyond rider check-in (standard ERPNext HRMS)",
	"Kitchen printer auto-routing per station (field exists; awaits hardware)",
	"Live provider calls for Foodpanda, WhatsApp, payment gateway, FBR — built as adapters in Mock mode until credentials arrive",
	"Docker deployment — not required (owner decision); production via bench",
	"Table reservations and online pre-booking",
], size=13)

# 4. nine steps
s = slide("One system, nine automatic steps", "2 · System overview")
steps = [("1", "Warehouse\ninventory"), ("2", "Kitchen\ninventory"), ("3", "Order\nplaced"), ("4", "Kitchen\ndisplay"), ("5", "Rider\nnotified"),
		 ("6", "Delivered"), ("7", "Review"), ("8", "Stock\nauto-update"), ("9", "Reports")]
x0, y0, bw, gap = Inches(0.55), Inches(1.6), Inches(1.27), Inches(0.12)
for i, (n, lbl) in enumerate(steps):
	x = x0 + i * (bw + gap)
	sh = box(s, x, y0, bw, Inches(0.7), "", fill=TEAL if i % 2 == 0 else TEAL_D, line=None, shape=MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON)
	text(s, x + Inches(0.15), y0 + Inches(0.12), bw - Inches(0.3), Inches(0.45), n, size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
	text(s, x - Inches(0.05), y0 + Inches(0.75), bw + Inches(0.1), Inches(0.55), lbl, size=11, bold=True, color=TEAL_D, align=PP_ALIGN.CENTER)
detail = [
	("1–2", "Central Store and each branch kitchen are ERPNext warehouses; ingredients carry reorder levels on both."),
	("3", "Every channel creates a Restaurant Order with the channel, type (dine-in, takeaway, delivery, pickup) and lines routed to stations."),
	("4", "The order appears on the KDS instantly (Socket.IO); cooks bump lines, the card moves Placed → Preparing → Ready."),
	("5", "At Ready, a delivery is offered to the nearest free rider (or assigned from the KDS / dispatch map)."),
	("6", "Rider picks up, delivers with the customer's 4-digit code, collects cash; the invoice is made at delivery."),
	("7", "The customer rates the order from the tracking link; low ratings alert managers."),
	("8", "When the food leaves the kitchen, the dish recipes (BOM) are issued from the kitchen warehouse; below reorder level ERPNext raises a Material Request."),
	("9", "Sales, kitchen times, menu engineering, customers, variance — live dashboard and reports."),
]
for i, (n, d) in enumerate(detail):
	y = Inches(3.15) + i * Inches(0.47)
	box(s, Inches(0.55), y, Inches(0.7), Inches(0.38), n, fill=TEAL_L, line=None, size=12, bold=True, color=TEAL_D)
	text(s, Inches(1.4), y, Inches(11.4), Inches(0.4), d, size=13, anchor=MSO_ANCHOR.MIDDLE)

# 5. users & roles
s = slide("Users, devices and access", "1 · Users")
table(s, Inches(0.55), Inches(1.3), Inches(12.2), [
	["User", "Device / screen", "Main tasks", "Access rule"],
	["Cashier", "Counter PC — POS terminal", "Open / close shift, sell, hold, return*, cash in/out, receive stock", "On the counter's POS Profile user list"],
	["Customer at counter", "2nd monitor — Customer display", "Sees items, offers, total, change and FBR QR as the cashier rings up", "None (mirrors the terminal on the same PC)"],
	["Waiter / host", "Tablet — Table map", "Seat guests, send rounds to kitchen, move table, settle bill", "Restaurant Order write"],
	["Kitchen staff", "Wall screen — KDS", "See tickets by station, bump lines, recall, assign rider", "Restaurant Order read"],
	["Dispatcher", "PC — Dispatch map", "Phone orders, assign riders, collect rider cash", "Restaurant Order write"],
	["Rider", "Phone — Rider app", "Go on duty, accept / reject, pick up, deliver (code, photo, cash)", "Own deliveries only"],
	["Customer", "Phone / web — Ordering site, tracking", "Order, pay, track rider, rate order, earn points", "Public, rate limited"],
	["Walk-in guest", "Tablet — Kiosk", "Self-order; pay at counter", "Public, rate limited"],
	["Manager / owner", "Any — Dashboard, reports", "Approve big discounts & returns, reviews, campaigns, settings", "System / Sales / Accounts Manager"],
], [1.6, 2.6, 4.9, 3.1], size=11, row_h=0.46)
text(s, Inches(0.55), Inches(6.45), Inches(12), Inches(0.4), "* Returns and cashier discounts above 10% need a manager's sign-in at the terminal.", size=11, color=MUTED, italic=True)

# 6. architecture
s = slide("System overview — how the parts fit", "2 · System overview")
layers = [
	("Users' screens (React, one app at /ros)", ["POS terminal", "Customer display", "Table map", "KDS", "Dispatch", "Rider app", "Ordering site", "Kiosk", "Owner dashboard"], BLUE_L, BLUE),
	("restaurant_os — custom Frappe app", ["Order service", "KDS & tables", "Riders & tracking", "Stock deduction", "Loyalty & reviews", "Analytics", "Approvals", "Health"], TEAL_L, TEAL_D),
	("ERPNext v16 + POS Awesome + HRMS", ["POS Invoice / Sales Invoice", "POS Shifts", "Offers & Coupons", "BOM / Stock", "Material Request", "Loyalty Program", "Employee Checkin", "Payment Entry"], BG, MUTED),
	("Integration adapters (Mock ↔ Live)", ["Foodpanda", "Payment gateway", "WhatsApp", "FBR"], ORANGE_L, ORANGE),
]
for i, (title_, parts, fill, line) in enumerate(layers):
	y = Inches(1.35) + i * Inches(1.35)
	box(s, Inches(0.55), y, Inches(12.2), Inches(1.15), fill=fill, line=line, lw=1.25)
	text(s, Inches(0.7), y + Inches(0.06), Inches(8), Inches(0.3), title_, size=12, bold=True, color=line)
	n = len(parts)
	pw = (Inches(11.9) - Inches(0.12) * (n - 1)) // n
	for j, p in enumerate(parts):
		box(s, Inches(0.7) + j * (pw + Inches(0.12)), y + Inches(0.42), pw, Inches(0.58), p, fill=WHITE, line=line, size=11, bold=True, color=INK)
text(s, Inches(0.55), Inches(6.75), Inches(12.2), Inches(0.35),
	 "Live updates travel over Socket.IO (KDS, tables, dispatch, rider location). Background jobs run on Frappe workers in production.",
	 size=11, color=MUTED, italic=True)

# 7. DFD notation
s = slide("Data Flow Diagram notation", "3 · DFD")
entity(s, Inches(0.8), Inches(1.6), Inches(2.4), Inches(0.9), "External entity")
text(s, Inches(3.4), Inches(1.6), Inches(9), Inches(0.9), "A person or outside system that sends data in or receives data out (customer, rider, Foodpanda, FBR).",
	 size=14, anchor=MSO_ANCHOR.MIDDLE)
process(s, Inches(0.8), Inches(2.85), Inches(2.4), Inches(0.9), "2.0", "Process")
text(s, Inches(3.4), Inches(2.85), Inches(9), Inches(0.9), "Work the system does that changes data. Numbered: 0 for the whole system (Level 0), 1.0 … 7.0 in Level 1.",
	 size=14, anchor=MSO_ANCHOR.MIDDLE)
store(s, Inches(0.8), Inches(4.15), Inches(2.4), Inches(0.75), "D2", "Data store")
text(s, Inches(3.4), Inches(4.1), Inches(9), Inches(0.9), "Data at rest — mapped to ERPNext / restaurant_os DocTypes (e.g. D2 Restaurant Orders).",
	 size=14, anchor=MSO_ANCHOR.MIDDLE)
arrow(s, Inches(0.9), Inches(5.6), Inches(3.1), Inches(5.6), color=INK, lw=1.75)
text(s, Inches(1.1), Inches(5.2), Inches(2), Inches(0.3), "data flow", size=11, color=MUTED, align=PP_ALIGN.CENTER)
text(s, Inches(3.4), Inches(5.3), Inches(9), Inches(0.7), "Data moving between entities, processes and stores; the label names the data.",
	 size=14, anchor=MSO_ANCHOR.MIDDLE)

# 8. DFD level 0 — central process, entities left and right, one in/out arrow pair each
s = slide("DFD Level 0 — context diagram", "3 · DFD")
P = box(s, Inches(5.17), Inches(1.3), Inches(3.0), Inches(5.35), "", fill=TEAL, line=TEAL_D, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=2)
text(s, Inches(5.17), Inches(3.25), Inches(3.0), Inches(0.5), "0", size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
text(s, Inches(5.17), Inches(3.7), Inches(3.0), Inches(0.5), "Restaurant OS", size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
text(s, Inches(5.27), Inches(4.25), Inches(2.8), Inches(0.9), "orders · kitchen · delivery\nbilling · stock · loyalty · reports", size=11, color=RGBColor(0xCC, 0xFB, 0xF1), align=PP_ALIGN.CENTER)
LEFT = [("Customer", "web · app · kiosk", "order, payment, rating", "quote, tracking, receipt"),
		("Cashier / Waiter", None, "sale, shift, cash, table rounds", "receipt, bill, Z-report"),
		("Kitchen staff", None, "bump, recall, rider pick", "tickets by station"),
		("Rider", None, "duty, GPS, delivered + cash", "delivery offer, route")]
RIGHT = [("Manager / Owner", None, "settings, approvals", "dashboard, alerts, reports"),
		 ("Foodpanda", None, "orders (webhook)", "status, availability"),
		 ("Payment gateway", None, "paid / failed", "payment session"),
		 ("FBR (tax)", None, "FBR invoice no. + QR", "invoice data"),
		 ("WhatsApp / SMS", None, "delivery reports", "review link, campaign")]


def pair(y_mid, ex_edge, px_edge, label_in, label_out, left):
	"""In-arrow (grey, entity → system) above, out-arrow (teal) below; labels sit above / below their arrow."""
	up, dn = y_mid - Inches(0.11), y_mid + Inches(0.11)
	a, b_ = (ex_edge, px_edge) if left else (ex_edge, px_edge)
	arrow(s, a, up, b_, up, color=MUTED, lw=1.5)
	arrow(s, b_, dn, a, dn, color=TEAL, lw=1.5)
	lx_ = min(a, b_) + Inches(0.1)
	wd = abs(b_ - a) - Inches(0.2)
	text(s, lx_, up - Inches(0.3), wd, Inches(0.26), label_in, size=10, color=INK, align=PP_ALIGN.CENTER)
	text(s, lx_, dn + Inches(0.03), wd, Inches(0.26), label_out, size=10, color=TEAL_D, align=PP_ALIGN.CENTER)


for k, (name, sub, li, lo) in enumerate(LEFT):
	y = Inches(1.45) + k * Inches(1.32)
	e = entity(s, Inches(0.45), y, Inches(2.0), Inches(0.8), name, sub)
	pair(cy(e), e.left + e.width, P.left, li, lo, True)
for k, (name, sub, li, lo) in enumerate(RIGHT):
	y = Inches(1.35) + k * Inches(1.08)
	e = entity(s, Inches(10.88), y, Inches(2.0), Inches(0.72), name, sub)
	pair(cy(e), e.left, P.left + P.width, li, lo, False)
text(s, Inches(0.45), Inches(6.78), Inches(9), Inches(0.28), "Grey arrows: data into the system  ·  Teal arrows: data out of the system  ·  Blue boxes: external entities",
	 size=10, color=MUTED, italic=True)
notes(s, "Level 0 shows Restaurant OS as a single process and every outside party that exchanges data with it. "
		 "Customers, staff and riders use the React screens; Foodpanda, the payment gateway, FBR and WhatsApp connect through adapters.")

# 9. DFD level 1 — grid: processes in rows 1/3/5, data stores in rows 2/4
s = slide("DFD Level 1", "3 · DFD — processes and data stores")
C = [Inches(0.3), Inches(2.55), Inches(5.65), Inches(8.75), Inches(11.25)]      # column lefts
PW, PH, SW, SH = Inches(2.05), Inches(0.78), Inches(2.05), Inches(0.42)
R1, R2, R3, R4, R5 = Inches(1.55), Inches(2.72), Inches(3.62), Inches(4.85), Inches(5.95)
PR = {
	"1": process(s, C[1], R1, PW, PH, "1.0", "Capture order"),
	"2": process(s, C[2], R1, PW, PH, "2.0", "Kitchen (KDS)"),
	"3": process(s, C[3], R1, PW, PH, "3.0", "Deliver & track"),
	"4": process(s, C[1], R3, PW, PH, "4.0", "Bill, pay, report tax"),
	"5": process(s, C[2], R3, PW, PH, "5.0", "Issue stock & reorder"),
	"6": process(s, C[3], R3, PW, PH, "6.0", "Reviews & loyalty"),
	"7": process(s, C[2], R5, PW, PH, "7.0", "Analyse & report"),
}
ST = {
	"D2": store(s, C[1], R2, SW, SH, "D2", "Restaurant Orders"),
	"D1": store(s, C[2], R2, SW, SH, "D1", "Menu & recipes"),
	"D6": store(s, C[3], R2, SW, SH, "D6", "Riders, assignments"),
	"D3": store(s, C[1], R4, SW, SH, "D3", "Invoices, payments"),
	"D4": store(s, C[2], R4 + Inches(0.3), SW, SH, "D4", "Stock ledger"),
	"D5": store(s, C[3], R4, SW, SH, "D5", "Customers, points, reviews"),
}
EW, EH = Inches(1.75), Inches(0.55)
EE = {
	"cust": entity(s, C[0], Inches(1.15), EW, EH, "Customer"),
	"staff": entity(s, C[0], Inches(1.9), EW, EH, "Cashier / Waiter"),
	"fp": entity(s, C[0], Inches(2.65), EW, EH, "Foodpanda"),
	"pay": entity(s, C[0], Inches(3.45), EW, EH, "Payment gateway"),
	"fbr": entity(s, C[0], Inches(4.2), EW, EH, "FBR"),
	"kit": entity(s, C[2] + Inches(0.15), Inches(0.62), EW, Inches(0.45), "Kitchen staff"),
	"rider": entity(s, C[4], Inches(1.5), EW, EH, "Rider"),
	"wa": entity(s, C[4], Inches(3.0), EW, EH, "WhatsApp"),
	"mgr": entity(s, C[3], R5 + Inches(0.1), EW + Inches(0.3), EH, "Manager / Owner"),
	"sup": entity(s, C[0], R5 + Inches(0.1), EW + Inches(0.3), EH, "Supplier / Store"),
}
r = lambda sh: sh.left + sh.width  # noqa: E731
b = lambda sh: sh.top + sh.height  # noqa: E731


def lab(x, y, s_, w=1.6, color=INK):
	t = text(s, x, y, Inches(w), Inches(0.22), s_, size=9, color=color, align=PP_ALIGN.LEFT)
	t.text_frame.word_wrap = False                            # one line; the box grows to fit
	t.fill.solid(); t.fill.fore_color.rgb = WHITE


# order capture
arrow(s, r(EE["cust"]), cy(EE["cust"]), PR["1"].left, PR["1"].top + Inches(0.15)); lab(r(EE["cust"]) + Inches(0.05), cy(EE["cust"]) + Inches(0.02), "cart", 0.5)
arrow(s, r(EE["staff"]), cy(EE["staff"]), PR["1"].left, cy(PR["1"]) + Inches(0.15)); lab(r(EE["staff"]) - Inches(0.05), cy(EE["staff"]) + Inches(0.02), "sale, rounds", 0.8)
arrow(s, r(EE["fp"]), cy(EE["fp"]), PR["1"].left + Inches(0.3), b(PR["1"])); lab(r(EE["fp"]) - Inches(0.02), cy(EE["fp"]) + Inches(0.02), "FP order", 0.6)
arrow(s, PR["3"].left + Inches(0.4), PR["3"].top, r(EE["cust"]), EE["cust"].top + Inches(0.12), color=TEAL)
lab(C[2] + Inches(1.9), Inches(1.15), "live tracking", 0.9, TEAL_D)
arrow(s, cx(PR["1"]), b(PR["1"]), cx(ST["D2"]), ST["D2"].top); lab(cx(PR["1"]) - Inches(0.95), b(PR["1"]) + Inches(0.08), "new order", 0.75)
arrow(s, ST["D1"].left, ST["D1"].top + Inches(0.05), r(PR["1"]) - Inches(0.15), b(PR["1"])); lab(C[2] - Inches(0.75), b(PR["1"]) + Inches(0.06), "prices", 0.5)
# kitchen
arrow(s, r(PR["1"]), cy(PR["1"]), PR["2"].left, cy(PR["2"])); lab(r(PR["1"]) + Inches(0.3), cy(PR["1"]) - Inches(0.28), "ticket", 0.6)
arrow(s, cx(EE["kit"]), b(EE["kit"]), cx(EE["kit"]), PR["2"].top, both=True); lab(cx(EE["kit"]) + Inches(0.1), b(EE["kit"]) + Inches(0.05), "tickets ⇄ bumps", 1.1)
arrow(s, PR["2"].left + Inches(0.15), b(PR["2"]), r(ST["D2"]) - Inches(0.2), ST["D2"].top); lab(PR["2"].left - Inches(0.2), b(PR["2"]) + Inches(0.22), "status", 0.5)
arrow(s, r(PR["2"]), cy(PR["2"]), PR["3"].left, cy(PR["3"])); lab(r(PR["2"]) + Inches(0.15), cy(PR["2"]) - Inches(0.28), "ready delivery", 0.95)
# delivery
arrow(s, r(PR["3"]), cy(PR["3"]) - Inches(0.12), EE["rider"].left, cy(EE["rider"]) - Inches(0.12), color=TEAL); lab(r(PR["3"]) + Inches(0.05), cy(PR["3"]) - Inches(0.4), "offer", 0.45, TEAL_D)
arrow(s, EE["rider"].left, cy(EE["rider"]) + Inches(0.12), r(PR["3"]), cy(PR["3"]) + Inches(0.12)); lab(r(PR["3"]) - Inches(0.05), cy(PR["3"]) + Inches(0.16), "GPS, code, cash", 1.0)
arrow(s, cx(PR["3"]), b(PR["3"]), cx(ST["D6"]), ST["D6"].top, both=True)
# billing
arrow(s, cx(ST["D2"]), b(ST["D2"]), cx(PR["4"]), PR["4"].top); lab(cx(ST["D2"]) + Inches(0.05), b(ST["D2"]) + Inches(0.12), "completed order", 1.05)
arrow(s, r(EE["pay"]), cy(EE["pay"]), PR["4"].left, cy(PR["4"]) - Inches(0.1)); lab(r(EE["pay"]) + Inches(0.02), cy(EE["pay"]) - Inches(0.28), "paid / failed", 0.75)
arrow(s, PR["4"].left, cy(PR["4"]) + Inches(0.22), r(EE["fbr"]), cy(EE["fbr"]), both=True); lab(r(EE["fbr"]) - Inches(0.05), cy(EE["fbr"]) + Inches(0.03), "invoice ⇄ FBR no.", 1.05)
arrow(s, cx(PR["4"]), b(PR["4"]), cx(ST["D3"]), ST["D3"].top); lab(cx(PR["4"]) + Inches(0.05), b(PR["4"]) + Inches(0.12), "invoice, payment", 1.05)
# stock
arrow(s, r(ST["D2"]) - Inches(0.3), b(ST["D2"]), PR["5"].left + Inches(0.3), PR["5"].top); lab(r(ST["D2"]) + Inches(0.05), b(ST["D2"]) + Inches(0.3), "food left kitchen", 1.1)
arrow(s, cx(ST["D1"]), b(ST["D1"]), cx(PR["5"]), PR["5"].top); lab(cx(ST["D1"]) + Inches(0.05), b(ST["D1"]) + Inches(0.08), "recipe", 0.5)
arrow(s, cx(PR["5"]), b(PR["5"]), cx(ST["D4"]), ST["D4"].top); lab(cx(PR["5"]) + Inches(0.05), b(PR["5"]) + Inches(0.32), "issue, wastage", 0.95)
arrow(s, ST["D4"].left, cy(ST["D4"]) + Inches(0.1), r(EE["sup"]), EE["sup"].top + Inches(0.1)); lab(C[0] + Inches(1.95), R5 - Inches(0.3), "reorder request", 1.05)
# loyalty & reviews
arrow(s, r(ST["D3"]), cy(ST["D3"]) - Inches(0.05), PR["6"].left + Inches(0.1), b(PR["6"])); lab(C[3] - Inches(0.95), R4 - Inches(0.36), "invoice → points", 1.05)
arrow(s, cx(PR["6"]), b(PR["6"]), cx(ST["D5"]), ST["D5"].top, both=True)
arrow(s, r(PR["6"]), cy(PR["6"]) - Inches(0.1), EE["wa"].left, cy(EE["wa"]), color=TEAL); lab(r(PR["6"]) + Inches(0.05), cy(PR["6"]) - Inches(0.48), "review link", 0.8, TEAL_D)
arrow(s, EE["wa"].left + Inches(0.3), b(EE["wa"]), r(PR["6"]), cy(PR["6"]) + Inches(0.18)); lab(r(PR["6"]) + Inches(0.15), cy(PR["6"]) + Inches(0.2), "rating", 0.5)
# analytics
arrow(s, cx(ST["D3"]), b(ST["D3"]), PR["7"].left, cy(PR["7"]) - Inches(0.1), color=MUTED)
arrow(s, cx(ST["D4"]), b(ST["D4"]), cx(PR["7"]), PR["7"].top, color=MUTED)
arrow(s, cx(ST["D5"]), b(ST["D5"]), r(PR["7"]), PR["7"].top + Inches(0.15), color=MUTED)
arrow(s, r(PR["7"]), cy(PR["7"]) + Inches(0.15), EE["mgr"].left, cy(EE["mgr"]), color=TEAL); lab(r(PR["7"]) + Inches(0.1), cy(PR["7"]) + Inches(0.2), "dashboard, reports", 1.15, TEAL_D)
text(s, Inches(9.0), Inches(6.85), Inches(4), Inches(0.25), "Orange: data stores (D1–D6) · Teal boxes: processes", size=9, color=MUTED, italic=True, align=PP_ALIGN.RIGHT)
notes(s, "Level 1 decomposes process 0 into seven processes. Orders (D2) drive the kitchen, delivery, billing and stock; invoices (D3) "
		 "drive loyalty points and FBR; stock (D4) and reviews (D5) feed analytics. Customer ratings arrive through the tracking page and the WhatsApp link.")

# 10. level 1 process descriptions
s = slide("Level 1 processes — inputs, outputs and where they live", "3 · DFD")
table(s, Inches(0.4), Inches(1.25), Inches(12.5), [
	["Process", "Inputs", "Outputs", "Implemented in"],
	["1.0 Capture order", "Cart / counter sale / table round / Foodpanda order / phone order", "Restaurant Order (D2), KDS ticket", "api.pos, api.online, api.tables, api.riders.new_delivery_order, aggregators.foodpanda"],
	["2.0 Kitchen (KDS)", "Tickets by station; cook bumps", "Line & order status, Ready event", "api.kds, services.orders (state machine)"],
	["3.0 Deliver & track", "Ready delivery, rider GPS, delivery code, cash", "Offer, assignment (D6), tracking page, COD into Rider Cash", "services.riders, api.riders"],
	["4.0 Bill, pay & report tax", "Completed order, payments, gateway result", "POS / Sales Invoice, Payment Entry (D3), FBR invoice no. + QR", "api.pos, services.delivery, services.payments, integrations.tax_authority.fbr"],
	["5.0 Issue stock & reorder", "Order left kitchen, recipe (D1), wastage", "Stock Entry (D4), sold-out flags, Material Request", "services.stock_deduction, services.availability, ERPNext reorder"],
	["6.0 Reviews & loyalty", "Invoice, rating, campaign", "Loyalty points, Order Review, alerts, messages (D5)", "services.loyalty, services.reviews, services.campaigns"],
	["7.0 Analyse & report", "D2–D5", "Owner dashboard, 5 script reports", "services.analytics, report/*"],
], [2.2, 3.5, 3.4, 3.4], size=10, row_h=0.62)

# 11. data stores
s = slide("Data stores — mapped to DocTypes", "3 · DFD")
table(s, Inches(0.4), Inches(1.25), Inches(12.5), [
	["Store", "Holds", "DocTypes (ERPNext standard / restaurant_os custom)"],
	["D1 Menu, recipes", "Dishes, prices, add-ons, recipes, stations", "Item, Item Price, BOM · Menu Modifier Group, Kitchen Station, Restaurant Branch"],
	["D2 Restaurant Orders", "Every order from every channel and its lines", "Restaurant Order (+ Item) · Restaurant Table, Floor Area"],
	["D3 Invoices & payments", "Bills, tenders, shifts, FBR numbers", "POS Invoice, Sales Invoice, Payment Entry, POS Opening / Closing Shift, POS Cash Movement, Integration Request"],
	["D4 Stock ledger", "Ingredient moves and counts", "Stock Entry, Stock Ledger Entry, Bin, Stock Reconciliation, Material Request · Wastage Log"],
	["D5 Customers, points, reviews", "Profiles, loyalty, ratings, campaigns", "Customer, Loyalty Program, Loyalty Point Entry · Order Review, Restaurant Campaign"],
	["D6 Riders & assignments", "Riders, duty, offers, GPS trail", "Employee, Employee Checkin · Rider, Delivery Assignment, Rider Location Log"],
], [2.4, 3.4, 6.7], size=11, row_h=0.62)

# 12–19 functional requirements
FR = [
	("Order capture — every channel in one queue", "5 · Functional requirements", [
		"**FR-1.1 ** POS counter sale creates the kitchen order on invoice submit (channel POS)",
		"**FR-1.2 ** Dine-in: seat a table, send rounds (marked “+ New Items” on the KDS), move table, settle one bill",
		"**FR-1.3 ** Website / installable app: menu with add-ons, sold-out hiding, delivery pin on map, delivery radius and minimum order",
		"**FR-1.4 ** Kiosk: self-order, idle reset; waits as a held sale until the cashier takes payment (channel Kiosk)",
		"**FR-1.5 ** Phone order: dispatcher enters customer, address and pin; cash on delivery",
		"**FR-1.6 ** Foodpanda: webhook (signed) or simulator; invoiced when delivered; Foodpanda price list",
		"**FR-1.7 ** Idempotent: a repeated post (client UUID / Foodpanda id) never creates a second order",
		"**FR-1.8 ** Quote = invoice: the customer sees exactly what will be billed (offers, tax, delivery fee, points)",
	]),
	("POS terminal, shifts and payments", "5 · Functional requirements", [
		"**FR-2.1 ** Open shift with counted float per payment mode; close with counted cash, variance, Z-report",
		"**FR-2.2 ** Sell by touch / barcode / keyboard; hold & resume; split tender Cash → Card → Bank Transfer; credit sale for named customers",
		"**FR-2.3 ** Offline: sales queue on the device and sync once, idempotently",
		"**FR-2.4 ** Sales tax 18% (on top of or included in price, by setting); FBR service fee line; FBR invoice no. and QR on receipt",
		"**FR-2.5 ** Returns (whole / part) refund through the chosen mode — manager approval",
		"**FR-2.6 ** Cashier discounts above 10% of the bill — manager approval; offers / coupons are pre-approved",
		"**FR-2.7 ** Cash in / out (expense, deposit, customer receipt), stock receiving at the counter",
		"**FR-2.8 ** Receipt: company logo, 80 / 58 mm or A4, auto-print option",
		"**FR-2.9 ** Customer-facing display (/ros/display): live cart with dish photos and offers, amount to pay / change, thank-you with FBR QR; menu & offers slideshow between customers",
	]),
	("Kitchen display and dine-in tables", "5 · Functional requirements", [
		"**FR-3.1 ** Board columns Placed → Preparing → Ready → Completed, live over Socket.IO, filter by station",
		"**FR-3.2 ** Bump a line or the whole card; recall; print ticket; age colours (amber / red minutes) and sound",
		"**FR-3.3 ** Lines no station makes (delivery fee) never hold an order back",
		"**FR-3.4 ** Delivery cards show the rider; “Assign rider” lists on-duty riders and off-duty riders (check in & assign)",
		"**FR-3.5 ** Table map by floor area: free / occupied / bill, guests, running total; settle produces one invoice for the sitting",
	]),
	("Delivery, riders and tracking", "5 · Functional requirements", [
		"**FR-4.1 ** Rider goes on / off duty = HR Employee Checkin (with GPS position)",
		"**FR-4.2 ** At Ready, the order is offered to the nearest (or least busy) free rider; reject → next rider; manual mode available",
		"**FR-4.3 ** Rider app: accept, navigate, pick up, deliver with the customer's 4-digit code or photo, collect cash",
		"**FR-4.4 ** Cash collected sits in Rider Cash until handed over at the counter (journal into the open shift)",
		"**FR-4.5 ** Dispatch map: riders and orders on OpenStreetMap, ETA, assign / reassign",
		"**FR-4.6 ** Customer tracking page: steps, ETA, rider on map, delivery code, pay-now for failed online payments",
		"**FR-4.7 ** Foodpanda-rider orders are tracked from Foodpanda's rider updates (own riders are not offered them)",
	]),
	("Online payment, offers, loyalty, reviews, campaigns", "5 · Functional requirements", [
		"**FR-5.1 ** Online payment through a gateway adapter; unpaid orders stay off the KDS; paid orders need nothing collected",
		"**FR-5.2 ** POS Offers apply on every channel (e.g. free Raita with Karahi); promo codes (POS Coupon) at checkout",
		"**FR-5.3 ** Saffron Rewards: 1 point per Rs 100, 1 point = Rs 1, 365-day expiry; customers with a mobile enrolled automatically",
		"**FR-5.4 ** Points redeemable at the counter and online (lower COD / online charge)",
		"**FR-5.5 ** After delivery, the tracking page asks for stars (overall, food, delivery, each dish) and a comment; WhatsApp link after a delay",
		"**FR-5.6 ** Rating at or below 2 stars notifies managers and shows on the dashboard",
		"**FR-5.7 ** Campaigns to Lapsed / New / Top spenders / Everyone with {name} {restaurant} {coupon} {link}; preview, then send",
	]),
	("Inventory, recipes and procurement", "5 · Functional requirements", [
		"**FR-6.1 ** Each dish has a BOM recipe; add-ons that are items become their own lines and recipes",
		"**FR-6.2 ** When food leaves the kitchen (Completed / Picked Up / Delivered) the recipe is issued from the branch kitchen warehouse",
		"**FR-6.3 ** Failed issues are logged and retried hourly; never block the sale",
		"**FR-6.4 ** Sold-out: an ingredient running out hides affected dishes on the web, kiosk and POS (and Foodpanda)",
		"**FR-6.5 ** Wastage Log with reason → Material Issue",
		"**FR-6.6 ** Variance report: theoretical (recipes) vs actual (counts) per ingredient, branch, period",
		"**FR-6.7 ** Reorder levels on Central Store (purchase) and kitchen (transfer) → Material Requests → Purchase Orders",
	]),
	("Owner dashboard and reports", "5 · Functional requirements", [
		"**FR-7.1 ** Live dashboard: today vs yesterday (net, orders, ticket, food cost %), hourly sales, channels, top items",
		"**FR-7.2 ** Open orders and late tickets, kitchen load by station, tables, riders, deliveries waiting",
		"**FR-7.3 ** Low stock, sold-out dishes, wastage today, alerts (late orders, waiting deliveries, low ratings)",
		"**FR-7.4 ** Reviews panel: 7-day average, recent reviews, low & unanswered",
		"**FR-7.5 ** Reports: Restaurant Sales Analysis (by channel, branch, item, category, hour, day part, cashier), Kitchen Performance, Menu Engineering, Customer Insights, Ingredient Variance — in React and the desk",
	]),
]
for title_, kicker, items in FR:
	s = slide(title_, kicker)
	bullets(s, Inches(0.55), Inches(1.35), Inches(12.2), Inches(5.6), items, size=15, gap=9)

# 20. state machine
s = slide("Order state machine", "6 · Rules")
st = {}
row1 = [("Placed", 0.5), ("Accepted", 2.25), ("Preparing", 4.0), ("Ready", 5.75)]
for name, x in row1:
	st[name] = box(s, Inches(x), Inches(2.0), Inches(1.5), Inches(0.7), name, fill=TEAL_L, line=TEAL_D, shape=MSO_SHAPE.ROUNDED_RECTANGLE, size=13, bold=True, color=TEAL_D)
for a_, b_ in (("Placed", "Accepted"), ("Accepted", "Preparing"), ("Preparing", "Ready")):
	arrow(s, r(st[a_]), cy(st[a_]), st[b_].left, cy(st[b_]), color=INK)
st["Completed"] = box(s, Inches(8.0), Inches(1.15), Inches(1.7), Inches(0.7), "Completed", fill=TEAL, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, size=13, bold=True, color=WHITE)
st["Picked Up"] = box(s, Inches(8.0), Inches(2.85), Inches(1.7), Inches(0.7), "Picked Up", fill=TEAL_L, line=TEAL_D, shape=MSO_SHAPE.ROUNDED_RECTANGLE, size=13, bold=True, color=TEAL_D)
st["Out"] = box(s, Inches(8.0), Inches(4.15), Inches(1.7), Inches(0.7), "Out for Delivery", fill=TEAL_L, line=TEAL_D, shape=MSO_SHAPE.ROUNDED_RECTANGLE, size=12, bold=True, color=TEAL_D)
st["Delivered"] = box(s, Inches(10.6), Inches(3.45), Inches(1.7), Inches(0.7), "Delivered", fill=TEAL, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, size=13, bold=True, color=WHITE)
arrow(s, r(st["Ready"]), cy(st["Ready"]) - i_(0.15), st["Completed"].left, cy(st["Completed"]), "dine-in / takeaway / pickup", lx=i_(6.3), ly=i_(1.2), size=9)
arrow(s, r(st["Ready"]), cy(st["Ready"]) + i_(0.15), st["Picked Up"].left, cy(st["Picked Up"]), "delivery", lx=i_(6.6), ly=i_(3.0), size=9)
arrow(s, cx(st["Picked Up"]), b(st["Picked Up"]), cx(st["Out"]), st["Out"].top, color=INK)
arrow(s, r(st["Out"]), cy(st["Out"]), st["Delivered"].left, cy(st["Delivered"]) + i_(0.15), "own rider", lx=i_(9.3), ly=i_(4.25), size=9)
arrow(s, r(st["Picked Up"]), cy(st["Picked Up"]), st["Delivered"].left, cy(st["Delivered"]) - i_(0.15), "Foodpanda rider", lx=i_(9.4), ly=i_(2.9), size=9)
aw = box(s, Inches(0.5), Inches(3.4), Inches(1.5), Inches(0.7), "Awaiting\nPayment", fill=ORANGE_L, line=ORANGE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, size=12, bold=True, color=ORANGE)
arrow(s, cx(aw), aw.top, cx(st["Placed"]), b(st["Placed"]), "paid online", lx=i_(1.3), ly=i_(2.9), size=9)
cn = box(s, Inches(3.0), Inches(3.4), Inches(2.6), Inches(0.7), "Cancelled / Rejected / Returned", fill=BG, line=MUTED, shape=MSO_SHAPE.ROUNDED_RECTANGLE, size=11, bold=True, color=MUTED)
table(s, Inches(0.5), Inches(4.75), Inches(7.2), [
	["Automatic step", "Fires at"],
	["Offer to a rider", "Ready (own-rider delivery)"],
	["Issue recipe ingredients", "Completed / Picked Up / Delivered"],
	["Make invoice (web, phone, Foodpanda)", "Delivered / Completed (counter: at sale)"],
	["Review link (WhatsApp)", "Delivery delay after Delivered / Completed"],
], [3.8, 3.4], size=11, row_h=0.38)

# 21. business rules
s = slide("Business rules (owner decisions)", "6 · Rules")
table(s, Inches(0.4), Inches(1.25), Inches(12.5), [
	["Rule", "Setting / value", "Where to change"],
	["Sales tax", "18%, added on top of the price (or included) — tick to apply", "Restaurant Settings → Tax"],
	["Payment modes", "Cash → Credit Card → Bank Transfer (+ Online Payment)", "Restaurant Settings → Payment modes"],
	["Invoice mode", "POS Invoice at the counter; merged into Sales Invoices at shift close", "Restaurant Settings"],
	["Stock deduction", "In the background when food leaves the kitchen", "Restaurant Settings → Stock"],
	["Foodpanda", "Invoiced at Delivered; Foodpanda price list; delivered by Foodpanda riders", "Restaurant Settings → Foodpanda"],
	["Manager approval", "Cashier discount > 10% of bill; every return / refund", "Restaurant Settings → Manager Approval"],
	["Delivery", "Fee Rs 150, minimum order Rs 500, radius per branch, OTP code optional", "Settings / Restaurant Branch"],
	["Loyalty", "1 pt per Rs 100, 1 pt = Rs 1, expires 365 days, online redemption on", "Loyalty Program “Saffron Rewards”"],
	["Low rating alert", "≤ 2 stars → managers notified", "Restaurant Settings → Reviews"],
	["Branches", "Multi-branch; Gulberg is the main branch", "Restaurant Branch"],
], [2.3, 6.2, 4.0], size=11, row_h=0.47)

# 22. integrations
s = slide("Integrations — adapters, Mock until credentials", "7 · Integrations")
table(s, Inches(0.4), Inches(1.25), Inches(12.5), [
	["Partner", "Direction", "Today", "To go Live"],
	["Foodpanda (Delivery Hero)", "Orders in (webhook); status & availability out", "Mock — simulator + signed webhook ready", "API access, webhook secret, product mapping"],
	["Payment gateway (Safepay / PayFast / JazzCash / Easypaisa)", "Hosted checkout out; paid / failed webhook in", "Mock — built-in test payment page", "Merchant account, keys, provider's checkout call"],
	["FBR (or PRA)", "Invoice data out; FBR invoice number in", "Mock — number and QR on receipts", "POS ID, token, PRAL URL, PCT code"],
	["WhatsApp / SMS", "Review links, campaigns out", "Mock — logged as Integration Requests", "Provider account, approved templates"],
	["Maps", "Tiles in the browser", "OpenStreetMap", "Paid tiles / Google key for heavy traffic"],
], [3.2, 3.2, 3.1, 3.0], size=11, row_h=0.7)
text(s, Inches(0.4), Inches(5.75), Inches(12.5), Inches(0.9),
	 "Every call is logged as an Integration Request with its payload. Live webhooks require an HMAC-SHA256 signature; in Mock mode the payment "
	 "webhook is closed and the Foodpanda webhook needs a signed-in manager.", size=12, color=MUTED)

# 23. NFR
s = slide("Non-functional requirements — measured", "8 · Quality")
table(s, Inches(0.4), Inches(1.25), Inches(7.3), [
	["Requirement", "Target", "Measured (p95)"],
	["Order visible on KDS", "< 1 s", "113 ms ✓"],
	["POS checkout (server)", "< 300 ms", "204 ms ✓"],
	["Online quote", "< 500 ms", "55 ms ✓"],
	["KDS board refresh", "< 1 s", "16 ms ✓"],
	["Tracking page", "< 300 ms", "2 ms ✓"],
], [3.0, 1.8, 2.5], size=12, row_h=0.5)
bullets(s, Inches(8.0), Inches(1.25), Inches(4.9), Inches(5.5), [
	"**Reliability: ** offline POS with idempotent sync; idempotent webhooks; failed stock / FBR jobs retried hourly",
	"**Observability: ** health endpoint, Frappe Error Log, Integration Request log",
	"**Deployability: ** bench (no Docker — owner decision); production via bench setup production",
	"**Load test: ** 40 rounds of counter sale + online order, rolled back (tests/load.py)",
], size=13, gap=10)

# 24. security
s = slide("Security and audit", "7 · Security")
bullets(s, Inches(0.55), Inches(1.35), Inches(6.0), Inches(5.6), [
	"**Server-side checks ** on all 94 API methods — the screen never decides access",
	"**Counter: ** only users on the POS Profile, or managers",
	"**Riders: ** only their own deliveries",
	"**Public pages: ** rate limited; prices recomputed on the server; orders reached by a random 12-character key",
	"**Manager approval ** with password check, wrong-password lockout (5 per 10 min)",
], size=14, gap=10)
bullets(s, Inches(6.9), Inches(1.35), Inches(6.0), Inches(5.6), [
	"**Audit trail: ** change history on orders, invoices, stock entries, reconciliations, offers, coupons, cash movements, payments, prices",
	"**Approvals noted ** on the invoice: who approved, which cashier",
	"**Signed webhooks ** (HMAC-SHA256) for payments and Foodpanda",
	"**Secrets ** stored as Password fields, never returned by the API",
], size=14, gap=10)

# 25. acceptance
s = slide("Acceptance criteria — definition of done", "9 · Acceptance")
text(s, Inches(0.55), Inches(1.3), Inches(12.2), Inches(0.8),
	 "A test order from each channel flows automatically through KDS → rider → delivered → review, deducts the right ingredients from the right "
	 "kitchen warehouse, triggers reorder below threshold, and shows up in every report.", size=15, color=MUTED)
table(s, Inches(0.55), Inches(2.35), Inches(12.2), [
	["Check", "Result"],
	["6 channels: POS, dine-in, web, kiosk, Foodpanda, phone", "Pass — tests/test_e2e.py"],
	["On the KDS immediately (< 1 s)", "Pass"],
	["Own-rider deliveries offered at Ready, delivered with code and cash", "Pass"],
	["Every order invoiced; every customer can rate; low rating alerts managers", "Pass"],
	["0.35 kg chicken per Chicken Tikka issued from Gulberg Kitchen, once per order", "Pass"],
	["Below reorder level → Material Request", "Pass"],
	["All channels in Sales by Channel; reviews on the dashboard", "Pass"],
	["Unit tests (counter, offers, KDS, tables, tax, FBR, riders, online, kiosk, payments, loyalty, reviews, approvals)", "36 pass"],
], [8.7, 3.5], size=12, row_h=0.44)

# 26. open items
s = slide("Open items and next steps", "9 · Next")
bullets(s, Inches(0.55), Inches(1.35), Inches(6.0), Inches(5.6), [
	"**Credentials to go live: ** FBR / PRA, Foodpanda, payment gateway, WhatsApp",
	"**Hardware: ** receipt printer and KDS screen models; kitchen printers for per-station auto-print",
	"**Decision: ** let cashiers pick manual POS Offers?",
	"**Production: ** HTTPS server (rider GPS), workers, backups",
], size=14, gap=12)
bullets(s, Inches(6.9), Inches(1.35), Inches(6.0), Inches(5.6), [
	"**Where hardware goes: ** receipt printer — POS terminal → Printer; customer display — terminal → Customer display (2nd monitor); KDS — /ros/kds on the kitchen screen, pick its station; kiosk — /ros/kiosk full-screen; riders — /ros/rider on their phones",
	"**Training videos: ** one simulator video per operation (separate files)",
], size=14, gap=12)

prs.save(OUT)
print(f"saved {OUT} — {len(prs.slides)} slides")
